#!/usr/bin/env python3
"""Kurs ishi Word (.docx) yig'uvchisi.

Ishlatish:
    python3 yigish.py <talaba_papkasi>            # .docx yig'adi, PDF orqali mundarija raqamlarini
                                                   # topadi va betlar to'lishi haqida hisobot beradi
    python3 yigish.py <talaba_papkasi> --tez       # faqat .docx (mundarija raqamlari eskisicha)

Manba: <talaba_papkasi>/manba/malumot.md (anketa jadvali) va manba/matn/*.md (belgilash README.md da).
Natija: <talaba_papkasi>/<Familiya_Ism>.docx, tekshiruv PDF va rasmlar: manba/tekshiruv/.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

DASTUR = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(DASTUR)
LOGO = os.path.join(REPO, "shablon", "logo.png")

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

STANDART = {"Fakultet": "Fizika-matematika", "Kafedra": "Fizika",
            "Yo'nalish": "Mexanika va matematik modellashtirish"}

MATN_ENI = 17.0          # sm: 21 - 2.5 - 1.5
SHRIFT = "Times New Roman"
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5}


# ----------------------------------------------------------------------------- yordamchilar
def uz(matn):
    """O'zbekcha tutuq belgilarini bir xil qiladi: o‘ g‘ (U+2018), qolgani ’ (U+2019)."""
    matn = re.sub(r'"([^"]*)"', "\u201c\\1\u201d", matn)
    matn = re.sub(r"([oOgG])['‘’ʻʼ`]", "\\1\u2018", matn)
    return re.sub(r"['ʼ`]", "\u2019", matn)


def malumot_oqi(papka):
    yol = os.path.join(papka, "manba", "malumot.md")
    d = dict(STANDART)
    for qator in open(yol, encoding="utf-8"):
        m = re.match(r"\|\s*([^|]+?)\s*\|\s*([^|]*?)\s*\|", qator)
        if m and m.group(1) not in ("Maydon", "---"):
            qiymat = re.sub(r"\s*\(standart\)\s*$", "", m.group(2))
            if qiymat:
                d[m.group(1)] = qiymat
    for k in ("Talaba", "Guruh", "Mavzu", "Rahbar"):
        if not d.get(k):
            sys.exit(f"malumot.md da '{k}' yo'q")
    d.setdefault("Yil", "2026")
    return d


# ----------------------------------------------------------------------------- formulalar (LaTeX -> OMML)
class Formulalar:
    """LaTeX formulalarni pandoc orqali Word formulasiga (OMML) aylantiradi."""

    def __init__(self, keshyol):
        self.keshyol = keshyol
        self.kesh = json.load(open(keshyol, encoding="utf-8")) if os.path.exists(keshyol) else {}

    def tayyorla(self, royxat):
        yangi = [f for f in dict.fromkeys(royxat) if f not in self.kesh]
        if not yangi:
            return
        with tempfile.TemporaryDirectory() as t:
            md = os.path.join(t, "f.md")
            with open(md, "w", encoding="utf-8") as f:
                for i, lat in enumerate(yangi):
                    f.write(f"F{i} ${lat}$\n\n")
            out = os.path.join(t, "f.docx")
            subprocess.run(["pandoc", md, "-f", "markdown", "-o", out], check=True)
            shutil.unpack_archive(out, os.path.join(t, "x"), "zip")
            tree = etree.parse(os.path.join(t, "x", "word", "document.xml"))
            ns = {"w": W_NS, "m": M_NS}
            for p in tree.iter(f"{{{W_NS}}}p"):
                txt = "".join(p.xpath(".//w:t/text()", namespaces=ns)).strip()
                m = re.match(r"F(\d+)", txt)
                if not m:
                    continue
                om = p.find(f".//{{{M_NS}}}oMath")
                if om is None:
                    sys.exit(f"Formula aylanmadi: {yangi[int(m.group(1))]}")
                self.kesh[yangi[int(m.group(1))]] = etree.tostring(om, encoding="unicode")
        for lat in yangi:
            if lat not in self.kesh:
                sys.exit(f"Formula aylanmadi: {lat}")
        json.dump(self.kesh, open(self.keshyol, "w", encoding="utf-8"), ensure_ascii=False, indent=0)

    def element(self, lat, olcham=28):
        el = parse_xml(self.kesh[lat].replace("<m:oMath", f'<m:oMath xmlns:w="{W_NS}"', 1)
                       if "xmlns:w=" not in self.kesh[lat] else self.kesh[lat])
        for r in el.iter(f"{{{M_NS}}}r"):
            rpr = etree.SubElement(r, qn("w:rPr"))
            f = etree.SubElement(rpr, qn("w:rFonts"))
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                f.set(qn(a), "Cambria Math")
            etree.SubElement(rpr, qn("w:sz")).set(qn("w:val"), str(olcham))
            etree.SubElement(rpr, qn("w:szCs")).set(qn("w:val"), str(olcham))
            r.remove(rpr)
            mrpr = r.find(f"{{{M_NS}}}rPr")
            r.insert(1 if mrpr is not None else 0, rpr)
        for npr in el.iter(f"{{{M_NS}}}naryPr"):
            if npr.find(f"{{{M_NS}}}limLoc") is None:
                ll = etree.SubElement(npr, f"{{{M_NS}}}limLoc")
                ll.set(f"{{{M_NS}}}val", "undOvr")
        return el


# ----------------------------------------------------------------------------- matnni o'qish
INLINE = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*|@(?:eq|rasm):[A-Za-z0-9_]+)")


def bloklar_oqi(papka):
    """matn/*.md -> bloklar ro'yxati."""
    matn = os.path.join(papka, "manba", "matn")
    bloklar = []
    for fayl in sorted(os.listdir(matn)):
        if not fayl.endswith(".md"):
            continue
        for qator in open(os.path.join(matn, fayl), encoding="utf-8"):
            q = qator.rstrip("\n").strip()
            if not q or q.startswith("%"):
                continue
            if q.startswith("# "):
                bloklar.append(("bob", q[2:].strip()))
            elif q.startswith("## "):
                bloklar.append(("bolim", q[3:].strip()))
            elif q.startswith("$$"):
                m = re.match(r"\$\$(.+?)\$\$\s*(?:\{#([\w\-]+)\})?\s*$", q)
                if not m:
                    sys.exit(f"Formula qatori noto'g'ri: {q}")
                bloklar.append(("formula", m.group(1).strip(), m.group(2)))
            elif q.startswith("!rasm"):
                qism = [s.strip() for s in q[5:].split("|")]
                bloklar.append(("rasm", qism[0], qism[1], float(qism[2]), qism[3]))
            elif q.startswith("|"):
                bloklar.append(("xatboshisiz", q[1:].strip()))
            else:
                bloklar.append(("matn", q))
    return bloklar


def raqamla(bloklar):
    """Formula va rasm raqamlari (havolalar uchun)."""
    eq, rasm = {}, {}
    n, bob, rn = 0, 0, 0
    for b in bloklar:
        if b[0] == "bob":
            bob = ROMAN.get(b[1].split(".")[0].strip(), bob + 1)
            rn = 0
        elif b[0] == "formula":
            n += 1
            if b[2]:
                if b[2] in eq:
                    sys.exit(f"Formula nomi takrorlangan: {b[2]}")
                eq[b[2]] = str(n)
        elif b[0] == "rasm":
            rn += 1
            rasm[b[1]] = f"{bob}.{rn}"
    return eq, rasm


# ----------------------------------------------------------------------------- Word qurish
class Quruvchi:
    def __init__(self, papka, info, formulalar, toc_betlar):
        self.papka, self.info, self.f, self.toc_betlar = papka, info, formulalar, toc_betlar
        self.doc = Document()
        self._uslublar()

    # --- umumiy sozlamalar
    def _uslublar(self):
        st = self.doc.styles["Normal"]
        st.font.name = SHRIFT
        st.font.size = Pt(14)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), SHRIFT)
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(1.25)
        pf.widow_control = False
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(2.0)
        s.header_distance, s.footer_distance = Cm(1.0), Cm(0.8)
        s.different_first_page_header_footer = True
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.first_line_indent = Cm(0)
        fp.paragraph_format.line_spacing = 1.0
        self._maydon(fp, "PAGE")
        # formulalar shrifti
        mp = parse_xml(f'<m:mathPr xmlns:m="{M_NS}"><m:mathFont m:val="Cambria Math"/><m:brkBin m:val="before"/>'
                       '<m:smallFrac m:val="0"/><m:dispDef/><m:lMargin m:val="0"/><m:rMargin m:val="0"/>'
                       '<m:defJc m:val="centerGroup"/><m:intLim m:val="subSup"/><m:naryLim m:val="undOvr"/></m:mathPr>')
        self.doc.settings.element.append(mp)
        # Word 2013+ rejimi ("Compatibility Mode" bo'lmasin)
        st_el = self.doc.settings.element
        compat = st_el.find(qn("w:compat"))
        if compat is None:
            compat = OxmlElement("w:compat")
            st_el.append(compat)
        for cs in compat.findall(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                compat.remove(cs)
        cs = OxmlElement("w:compatSetting")
        for k, v in (("w:name", "compatibilityMode"), ("w:uri", "http://schemas.microsoft.com/office/word"), ("w:val", "15")):
            cs.set(qn(k), v)
        compat.append(cs)

    def _maydon(self, p, kod):
        r = p.add_run()
        for tur, matn in (("begin", None), (None, kod), ("separate", None), (None, "1"), ("end", None)):
            if tur:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), tur)
                r._r.append(e)
            elif matn == kod:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = f" {kod} "
                r._r.append(e)
            else:
                e = OxmlElement("w:t")
                e.text = matn
                r._r.append(e)
        r.font.size = Pt(14)

    def para(self, align=None, indent=True, oraliq=None, oldin=0, keyin=0):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        if align is not None:
            pf.alignment = align
        if not indent:
            pf.first_line_indent = Cm(0)
        if oraliq is not None:
            pf.line_spacing = oraliq
        pf.space_before, pf.space_after = Pt(oldin), Pt(keyin)
        return p

    def run(self, p, matn, bold=False, italic=False, size=None, caps=False):
        r = p.add_run(uz(matn))
        r.bold, r.italic = bold or None, italic or None
        if size:
            r.font.size = Pt(size)
        if caps:
            r.text = r.text.upper()
        return r

    def qator(self, p, matn, size=None):
        """Matn ichidagi $formula$, **qalin**, *kursiv*, @eq:nom, @rasm:nom ni qo'shadi."""
        for bo in INLINE.split(matn):
            if not bo:
                continue
            if bo.startswith("$") and bo.endswith("$") and len(bo) > 1:
                p._p.append(self.f.element(bo[1:-1], olcham=2 * (size or 14)))
            elif bo.startswith("**"):
                self.qator_qalin(p, bo[2:-2], size)
            elif bo.startswith("*"):
                self.run(p, bo[1:-1], italic=True, size=size)
            elif bo.startswith("@eq:"):
                self.run(p, self.eq[bo[4:]], size=size)
            elif bo.startswith("@rasm:"):
                self.run(p, self.rasm[bo[6:]], size=size)
            else:
                self.run(p, bo, size=size)

    def qator_qalin(self, p, matn, size):
        for bo in re.split(r"(\$[^$]+\$)", matn):
            if bo.startswith("$"):
                p._p.append(self.f.element(bo[1:-1], olcham=2 * (size or 14)))
            elif bo:
                self.run(p, bo, bold=True, size=size)

    # --- muqova
    def muqova(self):
        i = self.info
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.15)
        self.run(p, "O'ZBEKISTON RESPUBLIKASI OLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI", bold=True)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.15, oldin=6)
        self.run(p, f"ABU RAYHON BERUNIY NOMIDAGI URGANCH DAVLAT UNIVERSITETI {i['Fakultet']} FAKULTETI "
                    f"{i['Kafedra']} KAFEDRASI".upper(), bold=True)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.0, oldin=24, keyin=24)
        p.add_run().add_picture(LOGO, width=Cm(4.2))
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.5)
        yonalish = i["Yo'nalish"]
        self.run(p, f"{i['Guruh']}-guruh {yonalish} guruhi talabasi {i['Talaba']}ning "
                    f"\u201c{i['Mavzu']}\u201d mavzusida yozgan")
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.0, oldin=30, keyin=36)
        self.run(p, "KURS ISHI", bold=True, size=72)
        for nom, ism in (("Bajardi:", i["Talaba"]), ("Tekshirdi:", i["Rahbar"]), ("Baholash:", "")):
            p = self.para(WD_ALIGN_PARAGRAPH.LEFT, False, 1.5, keyin=6)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(3.2))
            p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5))
            self.run(p, f"{nom}\t____________" + (f"\t{ism}" if ism else ""))
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.0, oldin=self.toc_betlar.get("_muqova", 60))
        self.run(p, f"Urganch-{i['Yil']}.", bold=True)

    # --- mundarija
    def mundarija(self, bloklar):
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, keyin=12)
        p.paragraph_format.page_break_before = True
        self.run(p, "MUNDARIJA", bold=True)
        for b in bloklar:
            if b[0] not in ("bob", "bolim"):
                continue
            p = self.para(WD_ALIGN_PARAGRAPH.LEFT, False)
            p.paragraph_format.left_indent = Cm(0 if b[0] == "bob" else 0.8)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(MATN_ENI), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            bet = self.toc_betlar.get(b[1], "0")
            self.run(p, f"{b[1]}\t{bet}")

    # --- asosiy matn
    def asosiy(self, bloklar):
        self.eq, self.rasm = raqamla(bloklar)
        n = 0
        for b in bloklar:
            t = b[0]
            if t == "bob":
                p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, keyin=12)
                p.paragraph_format.page_break_before = True
                p.paragraph_format.keep_with_next = True
                self.run(p, b[1].upper(), bold=True)
            elif t == "bolim":
                p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, oldin=6, keyin=6)
                p.paragraph_format.keep_with_next = True
                self.run(p, b[1], bold=True)
            elif t == "matn":
                self.qator(self.para(), b[1])
            elif t == "xatboshisiz":
                self.qator(self.para(indent=False), b[1])
            elif t == "formula":
                n += 1
                p = self.para(WD_ALIGN_PARAGRAPH.LEFT, False)
                ts = p.paragraph_format.tab_stops
                ts.add_tab_stop(Cm(MATN_ENI / 2), WD_TAB_ALIGNMENT.CENTER)
                ts.add_tab_stop(Cm(MATN_ENI), WD_TAB_ALIGNMENT.RIGHT)
                p.paragraph_format.keep_together = True
                p.add_run("\t")
                p._p.append(self.f.element(b[1]))
                p.add_run(f"\t({n})")
            elif t == "rasm":
                p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.0, oldin=6, keyin=4)
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(os.path.join(self.papka, "manba", "chizmalar", b[2]), width=Cm(b[3]))
                p = self.para(WD_ALIGN_PARAGRAPH.CENTER, False, 1.0, keyin=8)
                self.run(p, f"{self.rasm[b[1]]}-rasm.", bold=True, size=12)
                self.qator(p, " " + b[4], size=12)

    def saqla(self, yol):
        self.doc.save(yol)


# ----------------------------------------------------------------------------- PDF tekshiruvi
def pdfga(docx, papka):
    os.makedirs(papka, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", papka, docx],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.join(papka, os.path.splitext(os.path.basename(docx))[0] + ".pdf")


def norm(s):
    return re.sub(r"[^0-9a-zа-я]", "", uz(s).lower().replace("\u2018", "").replace("\u2019", ""))


def tahlil(pdf, bloklar):
    import pdfplumber
    natija = {"betlar": 0, "sarlavhalar": {}, "bosh_joy": []}
    with pdfplumber.open(pdf) as d:
        natija["betlar"] = len(d.pages)
        matnlar = [norm(p.extract_text() or "") for p in d.pages]
        boshla = 2
        for b in bloklar:
            if b[0] not in ("bob", "bolim"):
                continue
            k = norm(b[1])
            for i in range(boshla, len(matnlar)):
                if k in matnlar[i]:
                    natija["sarlavhalar"][b[1]] = i + 1
                    boshla = i
                    break
        pastki = d.pages[0].height - 2.0 / 2.54 * 72
        for i, p in enumerate(d.pages):
            obj = [o["bottom"] for o in p.chars + p.images + p.rects + p.lines + p.curves
                   if o["bottom"] < pastki + 3]
            eng = max(obj) if obj else 0
            natija["bosh_joy"].append(round(pastki - eng, 1))
    return natija


# ----------------------------------------------------------------------------- asosiy
def yigish(papka, toc_betlar):
    info = malumot_oqi(papka)
    bloklar = bloklar_oqi(papka)
    tekshir = os.path.join(papka, "manba", "tekshiruv")
    os.makedirs(tekshir, exist_ok=True)
    f = Formulalar(os.path.join(tekshir, "formulalar_kesh.json"))
    lat = [b[1] for b in bloklar if b[0] == "formula"]
    for b in bloklar:
        if b[0] in ("matn", "xatboshisiz", "rasm"):
            lat += [m[1:-1] for m in re.findall(r"\$[^$]+\$", b[-1])]
    f.tayyorla(lat)
    q = Quruvchi(papka, info, f, toc_betlar)
    cp = q.doc.core_properties
    cp.author = info["Talaba"]
    cp.last_modified_by = info["Talaba"]
    cp.title = info["Mavzu"]
    cp.comments = ""
    q.muqova()
    q.mundarija(bloklar)
    q.asosiy(bloklar)
    nom = re.sub(r"\s+", "_", info["Talaba"].strip()).replace("'", "")
    yol = os.path.join(papka, nom + ".docx")
    q.saqla(yol)
    return yol, bloklar, tekshir


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    papka = os.path.abspath(sys.argv[1])
    tez = "--tez" in sys.argv
    toc_yol = os.path.join(papka, "manba", "tekshiruv", "mundarija.json")
    toc = json.load(open(toc_yol, encoding="utf-8")) if os.path.exists(toc_yol) else {}
    yol, bloklar, tekshir = yigish(papka, toc)
    if tez:
        print("Yig'ildi:", yol)
        return
    for urinish in range(5):
        nat = tahlil(pdfga(yol, tekshir), bloklar)
        muq = toc.get("_muqova", 60)
        yangi_muq = max(0, round(muq + nat["bosh_joy"][0] - 4))
        if all(toc.get(k) == v for k, v in nat["sarlavhalar"].items()) and abs(yangi_muq - muq) <= 3:
            break
        toc = dict(nat["sarlavhalar"], _muqova=yangi_muq)
        json.dump(toc, open(toc_yol, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        yol, bloklar, tekshir = yigish(papka, toc)
    print("Yig'ildi:", yol)
    print("Betlar soni:", nat["betlar"])
    print("Mundarija:")
    for k, v in nat["sarlavhalar"].items():
        print(f"   {v:>3}  {k}")
    bob_bet = {v for k, v in nat["sarlavhalar"].items() if re.match(r"^[IVX]+\.", k)}
    print("Bet oxiridagi bo'sh joy (pt; >30 bo'lsa bet to'lmagan):")
    for i, g in enumerate(nat["bosh_joy"], 1):
        belgi = ""
        if g > 30 and i not in (1, 2, nat["betlar"]):
            belgi = "  <-- TO'LMAGAN" + ("  (keyingi bet yangi bob)" if (i + 1) in bob_bet else "")
        print(f"   {i:>3}: {g:6.1f}{belgi}")


if __name__ == "__main__":
    main()
